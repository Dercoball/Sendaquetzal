from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from manual_operativo_content import BRAND, DOC_DATE, DOC_SUBTITLE, DOC_TITLE, SECTIONS, TOC, VERSION


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output" / "manuales"
DOCX_PATH = OUT_DIR / "Manual_Operativo_Finaer_2026-09-16.docx"


BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F9FD"
GRAY = "D9D9D9"
DARK_GRAY = "404040"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_cell_text(cell, text, bold=False, color=None, size=8.7):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"

    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_text(cell, text, bold=True, color="FFFFFF", size=8.2)

    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_shading(cell, PALE_BLUE if ridx % 2 == 0 else "FFFFFF")
            set_cell_text(cell, text)
    doc.add_paragraph()


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.4)


def add_numbered(doc, number, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.4)


def add_paragraph(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.04
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(9.6)
    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 4)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(15 if level == 1 else 11.5)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    if level == 1:
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), LIGHT_BLUE)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)


def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.62)
        section.bottom_margin = Inches(0.62)
        section.left_margin = Inches(0.62)
        section.right_margin = Inches(0.62)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.6)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor.from_string(BLUE)

    return doc


def add_cover(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    footer = section.footer.paragraphs[0]
    footer.text = ""

    doc.add_paragraph()
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = brand.add_run(BRAND)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run(DOC_TITLE)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor.from_string("111111")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run(DOC_SUBTITLE)
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(DARK_GRAY)

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.LEFT
    info.style = "Table Grid"
    rows = [
        ("Empresa", BRAND),
        ("Documento", "Manual operativo"),
        ("Version", VERSION),
        ("Fecha", DOC_DATE),
    ]
    for ridx, row in enumerate(rows):
        for cidx, value in enumerate(row):
            cell = info.rows[ridx].cells[cidx]
            set_cell_shading(cell, BLUE if cidx == 0 else PALE_BLUE)
            set_cell_text(cell, value, bold=cidx == 0, color="FFFFFF" if cidx == 0 else DARK_GRAY, size=9)
    doc.add_paragraph()
    add_paragraph(
        doc,
        "Documento de uso interno para personal operativo y administrativo. "
        "El contenido debe actualizarse cuando cambien perfiles, permisos o reglas de operacion.",
    )

    doc.add_page_break()


def setup_headers(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    for sec in doc.sections:
        header = sec.header.paragraphs[0]
        header.text = ""
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(f"{BRAND} | Manual Operativo")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(DARK_GRAY)

        footer = sec.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Pagina ")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        add_page_number(footer)


def add_toc(doc):
    add_heading(doc, "Contenido", 1)
    for item in TOC:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.name = "Aptos"
        run.font.size = Pt(9.2)
    doc.add_page_break()


def add_sections(doc):
    for index, section in enumerate(SECTIONS):
        if index in (3, 6, 9, 12):
            doc.add_page_break()
        add_heading(doc, section["title"], 1)
        for block in section["blocks"]:
            if block["type"] == "paragraph":
                add_paragraph(doc, block["text"])
            elif block["type"] == "bullets":
                for item in block["items"]:
                    add_bullet(doc, item)
            elif block["type"] == "steps":
                for idx, item in enumerate(block["items"], 1):
                    add_numbered(doc, idx, item)
            elif block["type"] == "table":
                add_table(doc, block["headers"], block["rows"])


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = setup_document()
    add_cover(doc)
    setup_headers(doc)
    add_toc(doc)
    add_sections(doc)
    doc.core_properties.title = f"{DOC_TITLE} - {BRAND}"
    doc.core_properties.subject = "Manual operativo"
    doc.core_properties.author = BRAND
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()

