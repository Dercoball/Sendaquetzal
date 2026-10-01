from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output" / "manuales"
DOCX_PATH = OUT_DIR / "Manual_Sistema_Finaer_2026-09-16.docx"


BLUE = "1F4E79"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F9FD"
GRAY = "D9D9D9"
DARK_GRAY = "404040"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, size=None, bold=None, color=None):
    font = run.font
    font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True, color="000000")
    return p


def table(doc, headers, rows, widths=None, small=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.2 if small else 9.8, bold=True, color="FFFFFF")
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = widths[i]
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=8.7 if small else 9.5)
            set_cell_border(cell)
            set_cell_margins(cell)
            if idx % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string("000000")

    for name, size in [("Title", 25), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("000000")
        style.font.bold = True if "Heading" in name or name == "Title" else False

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = footer.add_run("Manual del Sistema Finaer | Página ")
    set_run_font(rr, size=8, color=DARK_GRAY)
    add_page_number(footer)


def add_cover(doc):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Manual del Sistema Finaer")
    set_run_font(r, size=26, bold=True, color="000000")

    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Instructivo operativo para usuarios, perfiles, permisos y módulos")
    set_run_font(r, size=12.5, color="000000")

    meta = [
        ("Sistema", "Finaer / Plataforma Web de Pagos"),
        ("Tipo de documento", "Manual de usuario y operación interna"),
        ("Fecha de elaboración", "16 de septiembre de 2026"),
        ("Alcance", "Acceso, usuarios, colaboradores, permisos, clientes, préstamos, pagos, cobranza, inversionistas, reportes, activos, comisiones y contenido web"),
        ("Audiencia", "Dirección, administración, gerencia de cobranza, recursos humanos, supervisores, promotores, gestores y personal de soporte"),
    ]
    table(doc, ["Campo", "Detalle"], meta, [Inches(1.6), Inches(5.7)])

    add_paragraph(
        doc,
        "Este manual explica dónde se realiza cada operación dentro del sistema y qué perfil debe utilizarse para ejecutarla. También documenta la diferencia entre puesto, tipo de usuario y permiso, porque el acceso a cada pantalla depende de la matriz de permisos asignada al perfil.",
    )
    add_paragraph(
        doc,
        "La referencia funcional se elaboró a partir de las pantallas, controladores, scripts y migraciones del proyecto revisado. Cuando una regla depende de configuración en base de datos, el manual indica el módulo donde se administra.",
    )
    doc.add_page_break()


def add_document_control(doc):
    add_heading(doc, "Control del documento")
    table(
        doc,
        ["Elemento", "Valor"],
        [
            ("Versión", "1.0"),
            ("Estado", "Documento operativo generado para revisión interna"),
            ("Sistema revisado", ".NET Framework 4.8 WebForms, SQL Server, Dapper/ADO.NET, jQuery, DataTables, jqWidgets, Bootstrap"),
            ("Ruta principal", "SourceCode/plataformawebpagos/MyERP"),
            ("Criterio de uso", "Usar como guía de operación. Validar nombres exactos de menús contra la base productiva si se agregaron permisos nuevos después de las migraciones revisadas."),
        ],
        [Inches(1.8), Inches(5.5)],
    )
    add_heading(doc, "Mapa de lectura", 2)
    add_bullets(
        doc,
        [
            "Para altas de personal y usuarios, revisar primero las secciones 3 y 4.",
            "Para cambios de roles o permisos, revisar la sección 5.",
            "Para operación diaria de créditos, pagos y cobranza, revisar las secciones 7 a 10.",
            "Para catálogos y configuración general, revisar la sección 14.",
            "Para pedir una actualización futura de este manual, usar el prompt del apéndice.",
        ],
    )
    doc.add_page_break()


def add_executive_summary(doc):
    add_heading(doc, "1 Resumen del sistema")
    add_paragraph(
        doc,
        "Finaer es un sistema web administrativo para una operación de microcréditos. Centraliza clientes y avales, solicitudes de préstamo, aprobación, pagos, cartera vencida, cobranza especializada, inversionistas, comisiones, activos, materiales, reportes y contenido público.",
    )
    add_paragraph(
        doc,
        "La aplicación usa autenticación propia en Login.aspx. Al iniciar sesión, carga en sesión el usuario, el empleado asociado, el tipo de usuario y la conexión de base de datos. Cada página protegida valida que la sesión exista y que el perfil tenga el permiso de esa pantalla.",
    )
    table(
        doc,
        ["Capa", "Qué contiene", "Archivos o carpetas principales"],
        [
            ("Presentación", "Pantallas WebForms, tablas, formularios, modales y navegación lateral.", "pages/*.aspx, js/app/*, vendor/*"),
            ("Lógica de pantalla", "Métodos WebMethod que reciben peticiones AJAX y aplican validaciones.", "pages/*/*.aspx.cs"),
            ("Dominio", "Clases de datos para clientes, empleados, préstamos, pagos, cobranza, inversión y catálogos.", "Clases/*.cs"),
            ("Datos", "Consultas SQL directas, Dapper, SqlClient y migraciones Flyway.", "flyway/migrations/*.sql"),
            ("Archivos", "Documentos, fotos, garantías y comprobantes.", "Uploads, FileUploader.ashx, documento"),
            ("Integraciones", "SMS/WhatsApp, correo, generación de PDF/Word y QR.", "Twilio, EASendMail, Syncfusion, iTextSharp, QRCoder"),
        ],
        [Inches(1.35), Inches(3.25), Inches(2.65)],
    )


def add_access_flow(doc):
    add_heading(doc, "2 Acceso y navegación")
    add_heading(doc, "2.1 Ingreso al sistema", 2)
    add_numbered(
        doc,
        [
            "Abrir la URL oficial del sistema y entrar a Login.aspx.",
            "Capturar Login y Contraseña.",
            "Presionar Entrar.",
            "Si las credenciales son correctas, el sistema registra la entrada en bitacora_login y redirige a la pagina_entrada configurada para el puesto.",
        ],
    )
    table(
        doc,
        ["Dato de sesión", "Uso operativo"],
        [
            ("usuario", "Identifica el login activo y se muestra en navegación."),
            ("id_usuario", "Se usa para validar permisos y registrar movimientos."),
            ("id_empleado", "Relaciona el usuario con colaborador, plaza y jerarquía."),
            ("id_tipo_usuario", "Define perfil de acceso, menú visible y alcance de datos."),
            ("path", "Selecciona la cadena de conexión, normalmente connbd."),
        ],
        [Inches(1.6), Inches(5.7)],
    )
    add_heading(doc, "2.2 Menú lateral", 2)
    add_paragraph(
        doc,
        "El menú se construye dinámicamente con los permisos del tipo de usuario. Si un perfil no tiene un permiso activo, la opción no aparece y la página también rechaza el acceso al validar TienePermisoPagina.",
    )
    table(
        doc,
        ["Grupo de menú", "Pantallas principales"],
        [
            ("Home", "Index.aspx o la página inicial configurada en posicion.pagina_entrada."),
            ("Clientes", "Customers/Customers.aspx y Customers/CustomerHistory.aspx."),
            ("Préstamos", "LoanRequest.aspx, LoanApprove.aspx, CreditIncreaseRequest.aspx."),
            ("Pagos", "Loans/Payments.aspx y Loans/PaymentsOverdue.aspx."),
            ("Cobranza", "Cobranza/Cartera.aspx, Cobranza/Folios.aspx, Cobranza/Reporte.aspx."),
            ("Inversionistas", "Investors.aspx, Investments.aspx, Utilities.aspx, InvestmentsDashboard.aspx."),
            ("Reportes", "Reports/ReportDefault.aspx."),
            ("Otros", "Assets.aspx, DeliveryMaterials.aspx, Calendar.aspx."),
            ("Configuración", "Colaboradores, puestos, usuarios, plazas, períodos, calendarios, días de paro, mensajes, categorías y tipos de cliente."),
            ("Web", "Preguntas frecuentes, acerca de, términos, privacidad y tutoriales."),
            ("Configuración de reglas", "Reglas de comisión y evaluación de empleados."),
        ],
        [Inches(1.55), Inches(5.75)],
    )


def add_roles(doc):
    add_heading(doc, "3 Roles, puestos y visibilidad")
    add_paragraph(
        doc,
        "En la operación actual el puesto del colaborador alimenta el tipo de usuario del acceso. En NewEmployee.aspx, el campo Puesto se guarda como id_posicion del empleado y como id_tipo_usuario del usuario. Por eso, al crear o editar un colaborador, el puesto debe seleccionarse con cuidado.",
    )
    table(
        doc,
        ["ID", "Perfil o puesto", "Uso operativo", "Regla importante"],
        [
            ("1", "Director", "Consulta directiva, reportes y operación de alto nivel.", "Si tiene plaza asignada, se acota a esa plaza; si no tiene plaza, puede ver todas."),
            ("2", "Coordinador", "Supervisión regional y reportes según permisos asignados.", "Puede aparecer en permisos históricos de cobranza; revisar matriz vigente."),
            ("3", "Ejecutivo", "Revisión final de créditos, seguimiento de equipo y cartera bajo su jerarquía.", "Ve créditos de supervisores y promotores relacionados."),
            ("4", "Supervisor", "Validación de campo, revisión de promotores y seguimiento de plaza.", "Ve colaboradores o créditos bajo su equipo."),
            ("5", "Promotor", "Captura de solicitudes, clientes y cobranza corriente.", "Ve su propia cartera y referencias de supervisor/ejecutivo cuando aplica."),
            ("6", "Superusuario", "Perfil técnico o administrativo superior.", "Se oculta en listados comunes para usuarios no superadmin."),
            ("7", "Gestor de Cobranza", "Gestión en campo de cartera vencida asignada.", "Sólo ve cuentas asignadas por cobranza_asignacion."),
            ("9", "Capturista", "Captura operativa o administrativa.", "Se trata como puesto administrativo en alta de colaborador."),
            ("10", "Gerente de Cobranza", "Control del módulo Cobranza.", "Sólo debe ver Cartera, Folios y Reporte semanal según V12."),
            ("11", "Recursos Humanos", "Alta y administración de colaboradores y puestos.", "Tiene permisos 8 y 10; no requiere plaza, supervisor ni ejecutivo."),
        ],
        [Inches(0.45), Inches(1.55), Inches(2.65), Inches(2.75)],
        small=True,
    )
    add_heading(doc, "3.1 Puestos administrativos", 2)
    add_paragraph(
        doc,
        "Director, Capturista, Gerente de Cobranza y Recursos Humanos son puestos administrativos. En la pantalla de alta de colaborador, Plaza, Supervisor y Ejecutivo son opcionales para estos perfiles. En puestos operativos, esos campos deben completarse para que funcionen filtros, reportes y visibilidad.",
    )
    add_heading(doc, "3.2 Visibilidad de información", 2)
    table(
        doc,
        ["Perfil", "Alcance de datos"],
        [
            ("Superusuario", "Sin filtro de visibilidad aplicado por UserVisibilityScope."),
            ("Gerente de Cobranza", "Ve cobranza de todas las plazas y puede filtrar por plaza."),
            ("Gestor de Cobranza", "Sólo ve préstamos asociados a su id_empleado en cobranza_asignacion activa."),
            ("Director", "Si tiene plaza, ve esa plaza; si no tiene plaza, ve todas."),
            ("Ejecutivo", "Ve su propia información y la de supervisores/promotores bajo su estructura."),
            ("Supervisor", "Ve su propia información, ejecutivo relacionado y promotores bajo su supervisión."),
            ("Promotor", "Ve su propia operación y referencias de su supervisor/ejecutivo."),
        ],
        [Inches(1.8), Inches(5.5)],
    )


def add_employee_user_sections(doc):
    add_heading(doc, "4 Alta de colaboradores y usuarios")
    add_heading(doc, "4.1 Alta integral de colaborador", 2)
    add_paragraph(
        doc,
        "Ruta recomendada: Configuración > Colaboradores > Nuevo. Esta opción abre Config/NewEmployee.aspx y permite capturar el expediente laboral completo junto con el usuario de acceso.",
    )
    add_numbered(
        doc,
        [
            "Entrar a Configuración > Colaboradores.",
            "Presionar Nuevo para abrir NewEmployee.aspx.",
            "Seleccionar Fecha de ingreso y Puesto. El Puesto define el tipo de usuario.",
            "Seleccionar Plaza, Supervisor y Ejecutivo cuando el puesto sea operativo.",
            "Capturar datos personales del colaborador: CURP, nombre, apellidos, teléfono, ocupación y domicilio.",
            "Capturar el Login y Contraseña inicial del usuario.",
            "Capturar datos del Aval cuando aplique y anexar documentación.",
            "Guardar. El sistema valida duplicados de CURP y Login antes de insertar o actualizar.",
        ],
    )
    table(
        doc,
        ["Dato", "Regla de captura"],
        [
            ("Puesto", "Obligatorio. Alimenta id_posicion del empleado e id_tipo_usuario del usuario."),
            ("Plaza", "Obligatoria para puestos operativos; opcional para administrativos."),
            ("Supervisor y Ejecutivo", "Obligatorios en flujo operativo para mantener jerarquía y filtros."),
            ("CURP colaborador", "Valida duplicado contra empleados activos no eliminados."),
            ("Login", "Debe ser único; si ya existe para otro colaborador, no permite guardar."),
            ("Contraseña", "Se guarda con hash MD5. En edición puede dejarse vacía para no cambiarla."),
            ("Aval", "Se registra como cliente aval y se relaciona por dirección/documentación cuando se captura."),
        ],
        [Inches(1.7), Inches(5.6)],
    )
    add_heading(doc, "4.2 Alta rápida o edición de acceso", 2)
    add_paragraph(
        doc,
        "Ruta: Configuración > Usuarios. Esta pantalla lista usuarios existentes y permite crear o editar un acceso, cambiar contraseña y vincular un empleado. Debe usarse cuando el colaborador ya existe o cuando sólo se ajusta el acceso.",
    )
    add_numbered(
        doc,
        [
            "Entrar a Configuración > Usuarios.",
            "Presionar Nuevo para crear un acceso o Editar para modificar uno existente.",
            "Capturar Nombre, Login, Email y Teléfono.",
            "Seleccionar Tipo usuario. Debe coincidir con el puesto operativo que corresponda.",
            "Vincular el Empleado usando el combo de empleados.",
            "Guardar.",
            "Para cambiar contraseña, usar el botón Contraseña de la tabla, escribir la nueva clave y guardar.",
        ],
    )
    add_paragraph(
        doc,
        "Importante: la pantalla de Usuarios tiene código histórico para permisos por usuario, pero la administración visible y vigente de permisos por perfil está en Puestos. Para cambios de permisos del rol, usar Configuración > Puestos.",
    )


def add_permissions(doc):
    add_heading(doc, "5 Permisos y cambio de roles")
    add_heading(doc, "5.1 Dónde se dan permisos", 2)
    add_paragraph(
        doc,
        "Ruta: Configuración > Puestos. En esta pantalla se listan los perfiles de posicion. Cada fila muestra el botón Permisos. Al abrirlo, el sistema muestra permisos disponibles y permisos seleccionados para ese perfil.",
    )
    add_numbered(
        doc,
        [
            "Entrar con un perfil que tenga acceso a Configuración > Puestos.",
            "Ubicar el puesto o perfil a modificar.",
            "Presionar Permisos.",
            "Mover permisos de la lista disponible a la lista seleccionada, o retirar los que no correspondan.",
            "Presionar Guardar. El sistema borra los permisos anteriores de ese perfil en permisos_tipo_usuario y vuelve a insertar la selección.",
            "Pedir al usuario que cierre sesión y vuelva a entrar para reconstruir el menú con los permisos actualizados.",
        ],
    )
    table(
        doc,
        ["Tabla", "Función"],
        [
            ("permisos", "Catálogo de pantallas, nombre interno, nombre visible y tipo de menú."),
            ("permisos_tipo_usuario", "Matriz principal: qué permiso tiene cada puesto o perfil."),
            ("usuario.id_tipo_usuario", "Perfil con el que se valida el acceso."),
            ("posicion.pagina_entrada", "Página inicial después del login."),
            ("permisos_usuario", "Tabla histórica de permisos directos por usuario; no es la ruta principal de navegación actual."),
        ],
        [Inches(2.0), Inches(5.3)],
    )
    add_heading(doc, "5.2 Tipos de permiso", 2)
    table(
        doc,
        ["Tipo", "Descripción"],
        [
            ("1", "Botón"),
            ("2", "Configuración"),
            ("3", "Panel"),
            ("4", "Web"),
            ("6", "Libre"),
            ("7", "Préstamos"),
            ("8", "Comisiones"),
            ("10", "Reportes"),
            ("11", "Inversionistas"),
            ("12", "Activos y otros"),
            ("30", "Cobranza"),
        ],
        [Inches(0.9), Inches(6.4)],
    )
    add_heading(doc, "5.3 Cambio de rol de un usuario", 2)
    add_numbered(
        doc,
        [
            "Si el usuario está ligado a un colaborador, entrar a Configuración > Colaboradores y editar el colaborador.",
            "Cambiar el Puesto cuando el cambio sea laboral; el puesto actualizará el tipo de usuario del acceso.",
            "Si sólo se corrige el acceso, entrar a Configuración > Usuarios y editar Tipo usuario.",
            "Verificar en Configuración > Puestos que el perfil nuevo tenga los permisos correctos.",
            "Cerrar sesión y volver a entrar con el usuario modificado para validar menú y acceso.",
        ],
    )
    add_heading(doc, "5.4 Caso Gerente de Cobranza", 2)
    add_paragraph(
        doc,
        "Para Gerente de Cobranza usar el perfil 10. La migración V12 restringe este perfil a los permisos 90, 92 y 93: Cartera, Folios y Reporte semanal. Si el gerente ve Clientes, Pagos, Reportes generales u otros menús, revisar permisos_tipo_usuario para id_tipo_usuario 10.",
    )
    table(
        doc,
        ["Permiso", "Pantalla", "Uso"],
        [
            ("90", "Cobranza/Cartera.aspx", "Cartera vencida, asignación de gestor, observaciones y registro de recuperación."),
            ("92", "Cobranza/Folios.aspx", "Asignación de blocks de folios a gestores."),
            ("93", "Cobranza/Reporte.aspx", "Reporte semanal de cobros, comisiones y resumen económico."),
        ],
        [Inches(0.8), Inches(2.4), Inches(4.1)],
    )
    add_heading(doc, "5.5 Caso Recursos Humanos", 2)
    add_paragraph(
        doc,
        "Para Recursos Humanos usar el perfil 11. La migración V13 asigna permisos 8 y 10: Colaboradores y Puestos. Es un perfil administrativo, por lo que no requiere plaza, supervisor ni ejecutivo en la alta.",
    )


def add_module_map(doc):
    add_heading(doc, "6 Mapa general de módulos")
    table(
        doc,
        ["Módulo", "Qué administra", "Pantallas clave"],
        [
            ("Configuración", "Catálogos base, usuarios, colaboradores, puestos, plazas, periodos, calendarios, mensajes, categorías y comisiones.", "pages/Config/*"),
            ("Clientes", "Clientes, avales, historial, estatus, condonación, demanda y reactivación.", "Customers.aspx, CustomerHistory.aspx"),
            ("Préstamos", "Solicitud, aprobación, garantías, aumento de crédito y seguimiento del crédito.", "LoanRequest.aspx, LoanApprove.aspx, CreditIncreaseRequest.aspx"),
            ("Pagos", "Cobranza corriente, abonos, pagos vencidos y tickets.", "Payments.aspx, PaymentsOverdue.aspx"),
            ("Cobranza", "Cartera de recuperación, asignación de gestores, folios y reporte semanal.", "Cartera.aspx, Folios.aspx, Reporte.aspx"),
            ("Inversionistas", "Alta de inversionistas, inversiones, retiros, utilidades y dashboard.", "Investors.aspx, Investments.aspx, Utilities.aspx, InvestmentsDashboard.aspx"),
            ("Reportes", "Reporte maestro de cierre y determinación de fondos.", "ReportDefault.aspx"),
            ("Activos y materiales", "Activos, entrega de materiales y calendario.", "Assets.aspx, DeliveryMaterials.aspx, Calendar.aspx"),
            ("Comisiones", "Reglas de evaluación y evaluación de colaboradores.", "ConfigRules.aspx, EmployeeEvaluation.aspx"),
            ("Web", "Contenido público editable: FAQ, acerca de, términos, privacidad y tutoriales.", "pages/Web/*"),
        ],
        [Inches(1.35), Inches(3.1), Inches(2.85)],
        small=True,
    )


def add_customers(doc):
    add_heading(doc, "7 Clientes y avales")
    add_paragraph(
        doc,
        "Ruta: Clientes. La pantalla Customers.aspx permite consultar clientes registrados, filtrar por plaza y jerarquía, revisar estatus y ejecutar acciones administrativas sobre el cliente.",
    )
    add_heading(doc, "7.1 Consulta de clientes", 2)
    add_numbered(
        doc,
        [
            "Entrar a Clientes.",
            "Seleccionar filtros de Plaza, Ejecutivo, Supervisor o Promotor cuando estén disponibles.",
            "Usar los filtros de la tabla para buscar por nombre, CURP, teléfono, monto o estatus.",
            "Abrir Historial para revisar préstamos anteriores, comportamiento de pago, avales y documentos.",
        ],
    )
    table(
        doc,
        ["Acción", "Cuándo se usa", "Efecto operativo"],
        [
            ("Historial", "Cuando se necesita validar comportamiento crediticio o documentos.", "Abre CustomerHistory.aspx."),
            ("Condonar", "Cuando Dirección autoriza castigar o perdonar saldo.", "Marca el antecedente y afecta el estatus del cliente/crédito según lógica de la pantalla."),
            ("Demanda", "Cuando se turna a jurídico.", "Bloquea al cliente para nuevas colocaciones y conserva historial."),
            ("Reactivar", "Cuando se habilita nuevamente a un cliente.", "Regresa al cliente a operación permitida."),
            ("Eliminar", "Cuando el registro fue capturado por error y no debe seguir visible.", "Borrado lógico."),
        ],
        [Inches(1.25), Inches(3.0), Inches(3.05)],
    )


def add_loans(doc):
    add_heading(doc, "8 Préstamos y aprobación")
    add_heading(doc, "8.1 Bandeja de solicitudes", 2)
    add_paragraph(
        doc,
        "Ruta: Préstamos > Solicitudes o LoanRequest.aspx. Esta bandeja permite buscar solicitudes por folio, cliente, monto, fechas, promotor y estatus. También abre la captura o revisión de una solicitud.",
    )
    add_numbered(
        doc,
        [
            "Entrar a Préstamos.",
            "Filtrar por folio, cliente, monto, fecha de solicitud, promotor o estatus.",
            "Presionar Nuevo Préstamo para abrir LoanApprove.aspx.",
            "Para continuar una solicitud existente, usar la acción de la fila.",
        ],
    )
    add_heading(doc, "8.2 Captura y aprobación", 2)
    add_paragraph(
        doc,
        "Ruta: LoanApprove.aspx. La pantalla concentra datos del préstamo, cliente, avales, revisión de supervisor, revisión de ejecutivo, garantías y resolución final.",
    )
    table(
        doc,
        ["Bloque", "Quién lo llena", "Contenido"],
        [
            ("Datos generales", "Promotor o capturista autorizado", "Fecha, tipo de cliente, monto, promotor y datos de historial."),
            ("Cliente", "Promotor o capturista", "CURP, nombre, teléfono, ocupación, domicilio, ubicación y documentos."),
            ("Aval 1", "Promotor o capturista", "Datos personales, domicilio, teléfono y documentos del aval."),
            ("Aval 2", "Cuando aplique", "Segundo respaldo para montos o riesgos especiales."),
            ("Supervisor", "Supervisor", "Confirmación de ubicación, nota de campo y garantías prendarias."),
            ("Ejecutivo", "Ejecutivo", "Dictamen final y autorización de crédito."),
        ],
        [Inches(1.35), Inches(1.8), Inches(4.15)],
    )
    add_heading(doc, "8.3 Aumento de crédito", 2)
    add_paragraph(
        doc,
        "Ruta: Préstamos > Solicitud de aumento o CreditIncreaseRequest.aspx. Se usa cuando el monto solicitado excede el límite normal del tipo de cliente o del historial. Supervisor y Ejecutivo registran la relación de aprobación.",
    )


def add_payments(doc):
    add_heading(doc, "9 Pagos corrientes y vencidos")
    add_heading(doc, "9.1 Pagos corrientes", 2)
    add_paragraph(
        doc,
        "Ruta: Pagos o Loans/Payments.aspx. Se usa para registrar pagos semanales de créditos activos, consultar saldo, aplicar abonos y generar evidencia o ticket cuando corresponda.",
    )
    add_numbered(
        doc,
        [
            "Buscar el préstamo por cliente, folio o filtros de cartera.",
            "Validar saldo, semana y monto a pagar.",
            "Capturar el pago o abono autorizado.",
            "Verificar que el saldo y estatus se actualicen antes de cerrar la operación.",
        ],
    )
    add_heading(doc, "9.2 Pagos vencidos", 2)
    add_paragraph(
        doc,
        "Ruta: Loans/PaymentsOverdue.aspx. Presenta cartera vencida antes o durante gestión de cobranza. La visibilidad se filtra por UserVisibilityScope: promotor, supervisor, ejecutivo, director, gerente o gestor según corresponda.",
    )


def add_collections(doc):
    add_heading(doc, "10 Cobranza especializada")
    add_paragraph(
        doc,
        "El módulo Cobranza se agregó por migraciones V8 a V12. Opera cartera vencida, gestores, folios y reportes semanales. La entrada a cobranza ocurre por semanas transcurridas desde la fecha de solicitud, con parámetro semana_caida_cobranza, por defecto 17.",
    )
    add_heading(doc, "10.1 Cartera de cobranza", 2)
    add_numbered(
        doc,
        [
            "Entrar a Cobranza > Cartera.",
            "Filtrar por plaza si el perfil tiene alcance de varias plazas.",
            "Abrir Ver para revisar cliente, avales, promotor, adeudo, fallas, multas, gasto de cobranza y total exigible.",
            "Si se tiene perfil de gerencia de cobranza, asignar o reasignar gestor.",
            "Guardar observaciones de seguimiento cuando la cuenta tenga gestor asignado.",
            "Registrar cobro con monto recuperado, canal de pago, folio y observación.",
        ],
    )
    table(
        doc,
        ["Regla", "Valor o comportamiento"],
        [
            ("Semana de caída", "Parámetro semana_caida_cobranza, por defecto 17."),
            ("Multa", "Parámetro monto_multa, por defecto $50 por falla."),
            ("Gasto de cobranza", "Parámetro porcentaje_gasto_cobranza, por defecto 10 por ciento del subtotal."),
            ("Bucket", "Se calcula con cobranza_regla_comision por rango de semanas."),
            ("Comisión gestor", "monto_recuperado por porcentaje de la regla del bucket."),
            ("Folio", "Si se captura y está disponible para el gestor, cambia a Usado."),
        ],
        [Inches(2.0), Inches(5.3)],
    )
    add_heading(doc, "10.2 Folios de cobranza", 2)
    add_paragraph(
        doc,
        "Ruta: Cobranza > Folios. Sólo gerencia de cobranza puede asignar blocks de folios. Se selecciona gestor, prefijo, desde y hasta. El sistema crea folios individuales con estatus Disponible.",
    )
    add_heading(doc, "10.3 Reporte semanal", 2)
    add_paragraph(
        doc,
        "Ruta: Cobranza > Reporte semanal. Permite revisar cobros por gestor, comisiones, no tangibles y diferencia de efectivo a entregar. Los canales no tangibles incluyen Formato de fallo, Depósito a cuenta oficial y Pago en oficina.",
    )
    table(
        doc,
        ["Perfil", "Puede hacer"],
        [
            ("Gerente de Cobranza", "Ver cartera, asignar gestor, asignar folios, registrar/consultar cobros y revisar reporte semanal."),
            ("Director", "También se considera gerencia de cobranza para asignaciones según UserVisibilityScope."),
            ("Gestor de Cobranza", "Ver sólo su cartera asignada y registrar recuperación según permisos vigentes."),
            ("Otros perfiles", "Sólo ven cobranza si la matriz de permisos lo permite; la V12 limita Gerente de Cobranza a cobranza."),
        ],
        [Inches(1.7), Inches(5.6)],
    )


def add_investors_reports_assets(doc):
    add_heading(doc, "11 Inversionistas")
    add_paragraph(
        doc,
        "El módulo Inversionistas administra inversionistas, inversiones, retiros, utilidades y dashboard directivo. Se usa para registrar capital aportado, rendimientos pactados, vencimientos y comprobantes de transferencia.",
    )
    table(
        doc,
        ["Pantalla", "Uso"],
        [
            ("Investors.aspx", "Alta, edición, suspensión o baja lógica de inversionistas."),
            ("Investments.aspx", "Pestañas para invertir y retirar. Calcula utilidad, vencimiento y total de retiro."),
            ("Utilities.aspx", "Estados o consulta de utilidades por inversionista."),
            ("InvestmentsDashboard.aspx", "Consulta directiva del capital fondeado y colocado."),
        ],
        [Inches(2.0), Inches(5.3)],
    )
    add_heading(doc, "12 Reportes y gastos")
    add_paragraph(
        doc,
        "Reports/ReportDefault.aspx concentra reportes de cartera, comisiones y determinación de fondos. Bills.aspx registra gastos o facturas de sucursal que alimentan el cierre semanal.",
    )
    table(
        doc,
        ["Operación", "Pantalla", "Resultado"],
        [
            ("Determinación de fondos", "ReportDefault.aspx", "Consolidado de promotores, fallas, efectivo, gastos y concentrado de caja."),
            ("Gastos de sucursal", "Bills.aspx", "Captura concepto, fecha, monto y comprobante para descontar en el corte."),
            ("Exportación", "DataTables", "Varias tablas permiten exportar CSV, XLS o PDF según pantalla."),
        ],
        [Inches(1.65), Inches(1.95), Inches(3.7)],
    )
    add_heading(doc, "13 Activos, materiales y calendario")
    table(
        doc,
        ["Pantalla", "Qué controla"],
        [
            ("Assets.aspx", "Activos fijos como motocicletas, equipo de cómputo, impresoras o bienes asignados."),
            ("DeliveryMaterials.aspx", "Entrega de materiales a plazas: contratos, pagarés, formatos, folletos y otros insumos."),
            ("Calendar.aspx", "Calendario de entregas, días de paro, días festivos o fechas de operación."),
        ],
        [Inches(2.0), Inches(5.3)],
    )


def add_config_commissions_web(doc):
    add_heading(doc, "14 Configuración general")
    add_paragraph(
        doc,
        "Configuración agrupa catálogos que alimentan el resto del sistema. Deben modificarse con cuidado porque impactan combos, filtros, cálculos y navegación.",
    )
    table(
        doc,
        ["Pantalla", "Función"],
        [
            ("Plazas.aspx", "Alta y edición de sucursales o plazas."),
            ("Positions.aspx", "Puestos, perfiles y permisos por perfil."),
            ("Puestos.aspx", "Catálogo de puestos si se usa separado de posiciones."),
            ("Usuarios.aspx", "Accesos, tipo usuario, empleado ligado y contraseña."),
            ("Employees.aspx", "Listado, filtros y edición de colaboradores."),
            ("NewEmployee.aspx", "Alta integral de colaborador, aval, documentos y usuario."),
            ("Periods.aspx", "Periodos operativos."),
            ("Calendars.aspx", "Calendarios base."),
            ("DaysOff.aspx", "Días de paro o no laborables."),
            ("Messages.aspx", "Mensajes usados en notificaciones Twilio."),
            ("CustomerTypes.aspx", "Tipos de cliente."),
            ("Categories.aspx", "Categorías generales."),
            ("CategoriesMaterials.aspx", "Categorías de materiales."),
            ("Commissions.aspx", "Catálogo de comisiones."),
        ],
        [Inches(2.0), Inches(5.3)],
        small=True,
    )
    add_heading(doc, "15 Comisiones y reglas")
    add_paragraph(
        doc,
        "El módulo de Configuración de reglas administra criterios para evaluar comisiones y desempeño de empleados. Las pantallas principales son Commissions/ConfigRules.aspx y Commissions/EmployeeEvaluation.aspx.",
    )
    add_heading(doc, "16 Contenido web")
    add_paragraph(
        doc,
        "El módulo Web administra contenido público: preguntas frecuentes, acerca de, términos y condiciones, aviso de privacidad y tutoriales. Estas pantallas usan contenido editable y se apoyan en TinyMCE para texto enriquecido.",
    )
    table(
        doc,
        ["Pantalla administrativa", "Contenido"],
        [
            ("WFAQ.aspx", "Preguntas frecuentes."),
            ("WAboutUs.aspx", "Acerca de la empresa."),
            ("WTermsAndConditions.aspx", "Términos y condiciones."),
            ("WNoticyOfPrivacy.aspx", "Aviso de privacidad."),
            ("WTutoriales.aspx", "Tutoriales."),
        ],
        [Inches(2.1), Inches(5.2)],
    )


def add_files_security(doc):
    add_heading(doc, "17 Documentos, archivos y seguridad")
    add_heading(doc, "17.1 Archivos y evidencias", 2)
    add_paragraph(
        doc,
        "El sistema maneja documentos de clientes, colaboradores, avales, garantías e inversiones. Los archivos se guardan en Uploads y también pueden guardarse como base64 o referencia en base de datos según el tipo de documento.",
    )
    table(
        doc,
        ["Tipo de archivo", "Uso"],
        [
            ("Fotografía", "Identificación visual de cliente, colaborador o aval."),
            ("INE frente y reverso", "Expediente documental obligatorio."),
            ("Comprobante domicilio", "Validación de domicilio."),
            ("Garantías", "Evidencia fotográfica de bienes prendarios."),
            ("Comprobantes inversión/retiro", "Soporte bancario de inversionistas."),
            ("Ticket de pago", "Plantilla ticket_pago_01.docx/pdf y generación de comprobantes."),
        ],
        [Inches(2.1), Inches(5.2)],
    )
    add_heading(doc, "17.2 Reglas de seguridad operativa", 2)
    add_bullets(
        doc,
        [
            "Usar siempre el botón Salir al terminar la operación.",
            "No compartir usuario o contraseña. Cada movimiento se relaciona con id_usuario.",
            "Después de cambiar permisos o rol, cerrar sesión y volver a iniciar.",
            "No modificar permisos de Superusuario salvo con autorización administrativa.",
            "Verificar plaza y jerarquía en colaboradores operativos para evitar que desaparezcan datos de filtros.",
            "Resguardar folios físicos de cobranza y conciliarlos contra Cobranza > Folios y Reporte semanal.",
        ],
    )


def add_quick_answers(doc):
    add_heading(doc, "18 Respuestas rápidas")
    table(
        doc,
        ["Pregunta", "Respuesta operativa"],
        [
            ("¿En qué módulo doy de alta usuarios?", "Configuración > Usuarios para alta rápida de acceso. Para alta completa de persona y usuario, Configuración > Colaboradores > Nuevo."),
            ("¿Qué perfil se usa para dar de alta usuarios?", "El perfil debe tener permiso a Configuración > Usuarios. Normalmente lo opera administración, superusuario o quien tenga el permiso 9 en la matriz."),
            ("¿Dónde se dan permisos?", "Configuración > Puestos, botón Permisos. La matriz se guarda en permisos_tipo_usuario."),
            ("¿Dónde cambio el rol de Gerente de Cobranza?", "En Configuración > Colaboradores se cambia el Puesto a Gerente de Cobranza, o en Configuración > Usuarios se corrige Tipo usuario a 10 si sólo se ajusta el acceso."),
            ("¿Qué debe ver Gerente de Cobranza?", "Sólo Cobranza > Cartera, Folios y Reporte semanal, permisos 90, 92 y 93."),
            ("¿Qué debe ver Recursos Humanos?", "Colaboradores y Puestos, permisos 8 y 10."),
            ("¿Por qué un usuario no ve una página?", "Puede faltarle el permiso al perfil, tener sesión vencida, no estar ligado a empleado o tener plaza/jerarquía incorrecta."),
            ("¿Por qué un gestor no ve cuentas?", "El gestor sólo ve cuentas activas asignadas en cobranza_asignacion."),
        ],
        [Inches(2.25), Inches(5.05)],
        small=True,
    )


def add_checklists(doc):
    add_heading(doc, "19 Checklists de operación")
    add_heading(doc, "19.1 Alta de colaborador con usuario", 2)
    add_bullets(
        doc,
        [
            "Puesto correcto seleccionado.",
            "Plaza, supervisor y ejecutivo completos si es perfil operativo.",
            "CURP del colaborador sin duplicado.",
            "Login único y sin espacios.",
            "Contraseña inicial capturada o política de cambio definida.",
            "Aval y documentos cargados cuando aplique.",
            "Usuario probó inicio de sesión y menú esperado.",
        ],
    )
    add_heading(doc, "19.2 Cambio de permisos de perfil", 2)
    add_bullets(
        doc,
        [
            "Identificar perfil exacto por ID y nombre.",
            "Respaldar o anotar permisos actuales antes de cambiar.",
            "Agregar sólo permisos necesarios.",
            "Guardar en Puestos > Permisos.",
            "Validar con usuario de prueba o sesión nueva.",
            "Confirmar que el menú visible coincide con responsabilidad del puesto.",
        ],
    )
    add_heading(doc, "19.3 Alta de Gerente de Cobranza", 2)
    add_bullets(
        doc,
        [
            "Crear o editar colaborador con Puesto 10 Gerente de Cobranza.",
            "Dejar Plaza, Supervisor y Ejecutivo como opcionales si no corresponde a una plaza específica.",
            "Verificar en Puestos que perfil 10 tenga permisos 90, 92 y 93 únicamente.",
            "Iniciar sesión y confirmar que sólo aparezca el módulo Cobranza.",
            "Probar asignación de cartera y folios con una cuenta controlada.",
        ],
    )


def add_prompt_appendix(doc):
    add_heading(doc, "20 Prompt para actualizar este manual")
    add_paragraph(
        doc,
        "Usa este prompt dentro de Codex cuando quieras regenerar o actualizar el manual después de cambios en el sistema:",
    )
    prompt = (
        "Revisa el proyecto Finaer completo y genera un manual profesional en Word y PDF para usuarios administrativos y operativos. "
        "El manual debe explicar todos los módulos del sistema, rutas de pantalla, roles, puestos, tipos de usuario, permisos por perfil, alta de usuarios, alta de colaboradores, cambio de roles, reglas de visibilidad, flujo de préstamos, pagos, cobranza, inversionistas, reportes, activos, comisiones y contenido web. "
        "Contrasta el contenido contra el código fuente, migraciones SQL, pantallas ASPX, code-behind C# y scripts JS. "
        "Incluye una matriz clara de perfiles, una guía paso a paso para dar de alta usuarios, dónde asignar permisos, cómo configurar Gerente de Cobranza y Recursos Humanos, y una sección de respuestas rápidas. "
        "El documento debe verse profesional, con portada, control del documento, tablas limpias, secciones numeradas y revisión visual antes de entregarlo."
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(prompt)
    set_run_font(r, size=9.5)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_document(doc)
    add_cover(doc)
    add_document_control(doc)
    add_executive_summary(doc)
    add_access_flow(doc)
    add_roles(doc)
    add_employee_user_sections(doc)
    add_permissions(doc)
    add_module_map(doc)
    add_customers(doc)
    add_loans(doc)
    add_payments(doc)
    add_collections(doc)
    add_investors_reports_assets(doc)
    add_config_commissions_web(doc)
    add_files_security(doc)
    add_quick_answers(doc)
    add_checklists(doc)
    add_prompt_appendix(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
